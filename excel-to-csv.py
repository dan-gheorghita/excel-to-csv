import os
import csv
from openpyxl import load_workbook

# Loop through every file in the current directory.
for excelFile in os.listdir('.'):
    # Skip non-Excel files.
    if not excelFile.endswith('.xlsx'):
        continue

    # Load the workbook.
    wb = load_workbook(filename=excelFile, read_only=True)
    baseName = os.path.splitext(excelFile)[0]

    # Loop through every sheet in the workbook.
    for sheetName in wb.sheetnames:
        ws = wb[sheetName]
        csvFileName = f"{baseName}_{sheetName}.csv"
        
        # Open a CSV file for writing.
        with open(csvFileName, 'w', newline='', encoding='utf-8') as csvFile:
            writer = csv.writer(csvFile)
            
            # Loop through every row in the sheet.
            for row in ws.iter_rows(values_only=True):
                # Write the row data to the CSV file.
                writer.writerow(list(row))
